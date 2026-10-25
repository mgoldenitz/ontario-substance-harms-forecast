"""Build excel/ontario_forecast.xlsx: the same holdout test and forecast in Excel formulas.

Excel's FORECAST.ETS fits additive exponential smoothing (AAA) with seasonality set to 4.
FORECAST.ETS.CONFINT gives the half-width of the prediction interval.
Input: data/processed/quarterly_on_ca.csv
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "excel"
OUT.mkdir(exist_ok=True)

SERIES = [
    ("Opioid deaths", "Opioids", "Deaths"),
    ("Stimulant deaths", "Stimulants", "Deaths"),
    ("Opioid hospitalizations", "Opioids", "Hospitalizations"),
    ("Stimulant hospitalizations", "Stimulants", "Hospitalizations"),
]
TRAIN_END = "2024 Q3"
FUTURE = ["2026 Q2", "2026 Q3", "2026 Q4", "2027 Q1", "2027 Q2"]

q = pd.read_csv(ROOT / "data" / "processed" / "quarterly_on_ca.csv")
q = q[q.Region == "Ontario"]
wide = q.pivot_table(index="Year_Quarter", columns=["Substance", "Source"], values="Value_num")
quarters = list(wide.index)

FONT = "Arial"
base = Font(name=FONT, size=10)
bold = Font(name=FONT, size=10, bold=True)
blue = Font(name=FONT, size=10, color="0000FF")
green = Font(name=FONT, size=10, color="008000")
title = Font(name=FONT, size=12, bold=True)
head_fill = PatternFill("solid", fgColor="D9E1F2")
thin = Border(bottom=Side(style="thin", color="808080"))
wrap = Alignment(wrap_text=True, vertical="top")


def q_date(label):
    year, qq = label.split(" Q")
    return f"DATE({year},{(int(qq) - 1) * 3 + 1},1)"


def header(ws, row, labels, widths=None):
    for c, text in enumerate(labels, start=1):
        cell = ws.cell(row, c, text)
        cell.font, cell.fill, cell.border = bold, head_fill, thin
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    if widths:
        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w


wb = Workbook()

# Notes
ws = wb.active
ws.title = "Notes"
ws.column_dimensions["A"].width = 110
notes = [
    ("Ontario substance harms forecast: Excel version", title),
    ("Source: Public Health Agency of Canada, Opioid- and stimulant-related harms in Canada, Health Infobase, "
     "SubstanceHarmsData.csv, downloaded 25 October 2026. Open Government Licence - Canada.", base),
    ("Data: Ontario quarterly counts (overall numbers). Opioid deaths and both hospitalization series run from "
     "Q1 2016 to Q1 2026; stimulant deaths start in Q1 2018.", base),
    ("Holdout: FORECAST.ETS is fit on quarters up to Q3 2024 and compared with the actual counts for Q4 2024 to Q1 2026. "
     "A seasonal naive forecast (the same quarter of the last training year) is the baseline.", base),
    ("Forecast: FORECAST.ETS is fit on every quarter and projects Q2 2026 to Q2 2027 with 80% and 95% intervals "
     "from FORECAST.ETS.CONFINT. Seasonality is set to 4 (quarterly).", base),
    ("Colours: blue numbers on the Data sheet are values typed in from the source file; black cells are formulas; "
     "green cells pull from another sheet.", base),
    ("Excel's ETS is a different model from the Python ARIMA, so the numbers will not match the Python results exactly.", base),
]
for i, (text, f) in enumerate(notes, start=1):
    c = ws.cell(i, 1, text)
    c.font, c.alignment = f, wrap

# Data
ws = wb.create_sheet("Data")
header(ws, 1, ["Quarter", "Quarter start"] + [s[0] for s in SERIES], [10, 13, 13, 13, 15, 15])
ws.row_dimensions[1].height = 30
for r, label in enumerate(quarters, start=2):
    ws.cell(r, 1, label).font = base
    d = ws.cell(r, 2, "=" + q_date(label))
    d.font, d.number_format = base, "yyyy-mm-dd"
    for c, (_, sub, src) in enumerate(SERIES, start=3):
        v = wide.loc[label, (sub, src)]
        if pd.notna(v):
            cell = ws.cell(r, c, int(v))
            cell.font = blue
ws["C1"].comment = Comment("Source: Health Infobase SubstanceHarmsData.csv, Region = Ontario, "
                           "Specific_Measure = Overall numbers, Time_Period = By quarter, Unit = Number.", "Source note")
ws.freeze_panes = "C2"
last_row = len(quarters) + 1
train_end_row = quarters.index(TRAIN_END) + 2
first_row = {}
for c, (name, sub, src) in enumerate(SERIES, start=3):
    s = wide[(sub, src)]
    first_row[name] = quarters.index(s.first_valid_index()) + 2

# Holdout
ws = wb.create_sheet("Holdout")
ws["A1"], ws["A1"].font = "Holdout test: fit to Q3 2024, test Q4 2024 to Q1 2026", title
test_quarters = quarters[train_end_row - 1:]
row0 = 3
for b, (name, _, _) in enumerate(SERIES):
    col_letter = get_column_letter(3 + b)
    top = row0 + b * (len(test_quarters) + 6)
    ws.cell(top, 1, name).font = bold
    header(ws, top + 1, ["Quarter", "Actual", "ETS forecast", "95% low", "95% high", "Seasonal naive",
                         "ETS abs error", "ETS abs % error", "Naive abs error", "Naive abs % error"],
           [10, 10, 12, 10, 10, 12, 12, 12, 12, 12])
    ws.row_dimensions[top + 1].height = 30
    yr = f"Data!${col_letter}${first_row[name]}:${col_letter}${train_end_row}"
    xr = f"Data!$B${first_row[name]}:$B${train_end_row}"
    for k, label in enumerate(test_quarters):
        r = top + 2 + k
        data_row = train_end_row + 1 + k
        naive_row = train_end_row - 3 + (k % 4)
        cells = [
            (1, label, base, None),
            (2, f"=Data!{col_letter}{data_row}", green, "#,##0"),
            (3, f"=_xlfn.FORECAST.ETS(Data!B{data_row},{yr},{xr},4)", base, "#,##0"),
            (4, f"=C{r}-_xlfn.FORECAST.ETS.CONFINT(Data!B{data_row},{yr},{xr},0.95,4)", base, "#,##0"),
            (5, f"=C{r}+_xlfn.FORECAST.ETS.CONFINT(Data!B{data_row},{yr},{xr},0.95,4)", base, "#,##0"),
            (6, f"=Data!{col_letter}{naive_row}", green, "#,##0"),
            (7, f"=ABS(C{r}-B{r})", base, "#,##0"),
            (8, f"=G{r}/B{r}", base, "0.0%"),
            (9, f"=ABS(F{r}-B{r})", base, "#,##0"),
            (10, f"=I{r}/B{r}", base, "0.0%"),
        ]
        for c, v, f, fmt in cells:
            cell = ws.cell(r, c, v)
            cell.font = f
            if fmt:
                cell.number_format = fmt
    first, last = top + 2, top + 1 + len(test_quarters)
    r = last + 1
    ws.cell(r, 1, "MAE / MAPE").font = bold
    for c, fmt in [(7, "#,##0.0"), (8, "0.0%"), (9, "#,##0.0"), (10, "0.0%")]:
        L = get_column_letter(c)
        cell = ws.cell(r, c, f"=AVERAGE({L}{first}:{L}{last})")
        cell.font, cell.number_format = bold, fmt
    ws.cell(r, 3, "Inside 95%:").font = bold
    cell = ws.cell(r, 4, f"=SUMPRODUCT((B{first}:B{last}>=D{first}:D{last})*(B{first}:B{last}<=E{first}:E{last}))"
                         f"&\" of {len(test_quarters)}\"")
    cell.font = bold

# Forecast
ws = wb.create_sheet("Forecast")
ws["A1"], ws["A1"].font = "Forecast Q2 2026 to Q2 2027, fit on all quarters", title
for b, (name, _, _) in enumerate(SERIES):
    col_letter = get_column_letter(3 + b)
    top = 3 + b * (len(FUTURE) + 4)
    ws.cell(top, 1, name).font = bold
    header(ws, top + 1, ["Quarter", "Quarter start", "Forecast", "80% low", "80% high", "95% low", "95% high"],
           [10, 13, 11, 10, 10, 10, 10])
    yr = f"Data!${col_letter}${first_row[name]}:${col_letter}${last_row}"
    xr = f"Data!$B${first_row[name]}:$B${last_row}"
    for k, label in enumerate(FUTURE):
        r = top + 2 + k
        ws.cell(r, 1, label).font = base
        d = ws.cell(r, 2, "=" + q_date(label))
        d.font, d.number_format = base, "yyyy-mm-dd"
        ci80 = f"_xlfn.FORECAST.ETS.CONFINT(B{r},{yr},{xr},0.8,4)"
        ci95 = f"_xlfn.FORECAST.ETS.CONFINT(B{r},{yr},{xr},0.95,4)"
        for c, v in [(3, f"=_xlfn.FORECAST.ETS(B{r},{yr},{xr},4)"),
                     (4, f"=C{r}-{ci80}"), (5, f"=C{r}+{ci80}"),
                     (6, f"=C{r}-{ci95}"), (7, f"=C{r}+{ci95}")]:
            cell = ws.cell(r, c, v)
            cell.font, cell.number_format = base, "#,##0"

wb.save(OUT / "ontario_forecast.xlsx")
print("Saved", OUT / "ontario_forecast.xlsx")
